import csv
import hashlib
from pathlib import Path

import pytest
from PIL import Image

from photocheck.domain import Check, Result, Status
from photocheck.export import export_excel
from photocheck.files import collect, inspect, load_image


def save(root, name="증명사진.png", fmt="PNG"):
    path = root / name
    Image.new("RGB", (120, 160), "white").save(path, format=fmt)
    return path


def test_candidate_collection(tmp_path):
    save(tmp_path)
    (tmp_path / "안내.txt").write_text("text")
    (tmp_path / "미지원.heic").write_bytes(b"bad")
    nested = tmp_path / "하위"
    nested.mkdir()
    save(nested, "동일이름.JPG", "JPEG")
    shallow = collect(tmp_path, False)
    deep = collect(tmp_path, True)
    assert len(shallow.paths) == 2 and shallow.excluded == 1
    assert len(deep.paths) == 3 and deep.excluded == 1
    assert collect(tmp_path, True, lambda: True).cancelled


def test_file_failure_types(tmp_path):
    empty = tmp_path / "빈.jpg"
    empty.touch()
    damaged = tmp_path / "손상.png"
    damaged.write_bytes(b"not an image")
    unsupported = save(tmp_path, "형식.gif", "GIF")
    for path in (empty, damaged, unsupported, tmp_path / "없음.jpg"):
        result = inspect(path, tmp_path)
        assert result.status == Status.ERROR
        assert result.direction_check == result.framing_check == Check.SKIP


def test_truncated_decode_and_mismatch(tmp_path):
    path = save(tmp_path)
    raw = path.read_bytes()
    path.write_bytes(raw[:len(raw) // 2])
    assert inspect(path, tmp_path).status == Status.ERROR
    mismatch = save(tmp_path, "확장자.jpg", "PNG")
    image, check, reasons = load_image(mismatch)
    try:
        assert check == Check.REVIEW
        assert "불일치" in reasons[0]
    finally:
        image.close()


@pytest.mark.parametrize("orientation", range(1, 9))
def test_exif_orientation_and_original_unchanged(tmp_path, orientation):
    path = tmp_path / "방향.jpg"
    exif = Image.Exif()
    exif[274] = orientation
    Image.new("RGB", (120, 160), "red").save(path, exif=exif)
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    image, check, reasons = load_image(path)
    assert image.size == ((160, 120) if orientation >= 5 else (120, 160))
    assert check == Check.PASS and not reasons
    image.close()
    inspect(path, tmp_path)
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


def test_animation_never_passes(tmp_path):
    path = tmp_path / "움직임.png"
    first = Image.new("RGB", (30, 40), "red")
    first.save(path, save_all=True, append_images=[Image.new("RGB", (30, 40), "blue")], duration=100, loop=0)
    image, check, reasons = load_image(path)
    assert check == Check.REVIEW and any("프레임" in s for s in reasons)
    image.close()


def test_independent_gates_and_analysis_exception(tmp_path):
    path = save(tmp_path)
    class Broken:
        def analyze(self, image, result):
            result.direction_check = Check.PASS
            raise RuntimeError("test failure")
    result = inspect(path, tmp_path, Broken())
    assert result.status == Status.REVIEW
    assert result.direction_check == result.framing_check == Check.REVIEW
    for field in ("file_check", "direction_check", "framing_check"):
        result = Result(path, path.name, file_check=Check.PASS, direction_check=Check.PASS, framing_check=Check.PASS)
        setattr(result, field, Check.REVIEW)
        assert result.finalize().status == Status.REVIEW


def test_excel_unicode_and_formula_like_filename_as_text(tmp_path):
    from openpyxl import load_workbook
    result = Result(tmp_path / "=사진.png", "=사진.png", reasons=["한글 사유"])
    path = tmp_path / "결과.xlsx"
    export_excel(path, [result])
    book = load_workbook(path)
    sheet = book.active
    rows = list(sheet.values)
    assert rows[1][0] == "=사진.png" and sheet['A2'].data_type == 's'
    assert rows[1][1] == "한글 사유" and sheet.max_column == 2
    assert not any("수동" in field or "최종" in field for field in rows[0])
    book.close()

import hashlib
import json
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from PIL import Image
from PySide6.QtCore import QEventLoop, QTimer
from PySide6.QtGui import QImage
from PySide6.QtWidgets import QApplication, QFileDialog, QLabel, QMessageBox

from photocheck.criteria import Criteria
from photocheck.domain import Result, Status
from photocheck.export import export_excel
from photocheck.photo_copy import copy_photos
from photocheck.ui import MainWindow


def read_rows(path):
    book = load_workbook(path)
    try:
        values = list(book.active.values)
        return [dict(zip(values[0],row)) for row in values[1:]]
    finally:
        book.close()


def make_results(root):
    results=[]
    for folder,status in [('첫 폴더',Status.REVIEW),('둘 폴더',Status.ERROR),('정상 폴더',Status.NORMAL)]:
        path = root/folder/'동일사진.jpg'
        path.parent.mkdir(parents=True,exist_ok=True)
        Image.new('RGB',(120,160),'white').save(path)
        results.append(Result(path,str(path.relative_to(root)),status=status,reasons=['확인 사유'],
                              tags={'글자 포함 의심'},criteria_settings=Criteria().snapshot()))
    return results


def test_excel_excludes_automatic_and_reclassified_normal(tmp_path):
    rows = make_results(tmp_path)
    before = rows[0].criteria_settings.copy(),rows[0].reasons.copy()
    rows[0].mark_normal()
    assert rows[0].automatic_status == Status.REVIEW and rows[0].manual_normal and rows[0].reviewed_at
    assert rows[0].finalize().status == Status.NORMAL
    assert (rows[0].criteria_settings,rows[0].reasons) == before
    path = tmp_path/'결과.xlsx'
    assert export_excel(path,rows) == 1
    actual = read_rows(path)
    assert actual == [{'파일명': rows[1].path.name, '판정사유': '확인 사유'}]
    with ZipFile(path) as archive:
        assert not any(b'<f>' in archive.read(name) for name in archive.namelist())
    book=load_workbook(path)
    assert book.active.freeze_panes == 'A2' and book.active.auto_filter.ref == 'A1:B2'
    assert book.active.max_column == 2
    book.close()


def test_all_normal_export_is_valid_header_only_workbook(tmp_path):
    path=tmp_path/'비어있는 결과.xlsx'
    assert export_excel(path,[Result(tmp_path/'a.jpg','a.jpg',status=Status.NORMAL)]) == 0
    assert read_rows(path) == []


@pytest.mark.parametrize('name',['=SUM(A1).jpg','+1.jpg','-1.jpg','@사진&<검수>.jpg'])
def test_formula_like_values_are_plain_excel_text(tmp_path,name):
    path=tmp_path/'결과.xlsx'
    export_excel(path,[Result(tmp_path/name,name,reasons=['=1+1 & <문자>'])])
    book=load_workbook(path)
    assert book.active['A2'].value == name and book.active['A2'].data_type == 's'
    assert book.active['B2'].value == '=1+1 & <문자>' and book.active['B2'].data_type == 's'
    book.close()


def test_copy_preserves_nested_duplicate_names_and_excludes_normal(tmp_path):
    rows=make_results(tmp_path/'원본')
    destination=tmp_path/'사본'; destination.mkdir()
    hashes = {r.path:hashlib.sha256(r.path.read_bytes()).hexdigest() for r in rows}
    report=copy_photos(rows,destination)
    assert report.copied == 2 and not report.errors
    for r in rows[:2]:
        assert (destination/r.relative_path).read_bytes() == r.path.read_bytes()
    assert not (destination/rows[2].relative_path).exists()
    assert hashes == {r.path:hashlib.sha256(r.path.read_bytes()).hexdigest() for r in rows}
    existing=(destination/rows[0].relative_path)
    existing.write_bytes(b'keep prior copy')
    assert len(copy_photos(rows,destination).skipped) == 2
    assert existing.read_bytes() == b'keep prior copy'
    rows[0].mark_normal()
    second=tmp_path/'다른 사본'; second.mkdir()
    assert copy_photos(rows,second).copied == 1
    assert not (second/rows[0].relative_path).exists()


def test_copy_missing_input_continues_and_same_source_is_skipped(tmp_path):
    rows=make_results(tmp_path/'원본')
    destination=tmp_path/'사본'; destination.mkdir()
    absent=Result(tmp_path/'missing.jpg','missing.jpg')
    report=copy_photos([absent,rows[0]],destination)
    assert report.copied == 1 and len(report.errors) == 1
    assert not (destination/'missing.jpg').exists()
    report=copy_photos(rows[:2],tmp_path/'원본')
    assert report.copied == 0 and len(report.skipped) == 2


@pytest.mark.parametrize('relative',['../outside.jpg','C:/outside.jpg'])
def test_copy_rejects_path_outside_chosen_folder(tmp_path,relative):
    destination=tmp_path/'사본'; destination.mkdir()
    source=tmp_path/'original.jpg'; source.write_bytes(b'original')
    report=copy_photos([Result(source,relative)],destination)
    assert report.copied == 0 and len(report.errors) == 1
    assert list(destination.iterdir()) == [] and source.read_bytes() == b'original'


def test_cancelled_copy_removes_incomplete_destination(tmp_path):
    destination=tmp_path/'사본'; destination.mkdir()
    source=tmp_path/'original.jpg'; source.write_bytes(b'x'*(3*1024*1024))
    calls=0
    def cancel():
        nonlocal calls
        calls+=1
        return calls >= 3
    report=copy_photos([Result(source,source.name)],destination,cancelled=cancel)
    assert report.cancelled and report.copied == 0 and not (destination/source.name).exists()
    assert source.stat().st_size == 3*1024*1024


def test_unavailable_destination_is_a_reported_failure(tmp_path):
    source=tmp_path/'original.jpg'; source.write_bytes(b'original')
    report=copy_photos([Result(source,source.name)],tmp_path/'없어진 폴더')
    assert report.copied == 0 and len(report.errors) == 1 and '폴더 접근 실패' in report.errors[0]
    assert source.read_bytes() == b'original'


def window_fixture(tmp_path):
    app=QApplication.instance() or QApplication([])
    window=MainWindow()
    rows=make_results(tmp_path/'원본')
    window.total=3
    window.received(rows)
    window.set_running(False)
    window.pool.waitForDone(); app.processEvents()
    return app,window,rows


def test_ui_columns_labels_removed_notices_and_normal_navigation(tmp_path):
    app,window,rows=window_fixture(tmp_path)
    assert window.model.headers == ['파일명','주요 사유'] and window.model.columnCount() == 2
    assert window.export.text() == '결과 저장' and window.copy.text() == '사진 복사'
    assert not any(label.text().startswith(('초기 보수','후보 확장자:')) for label in window.findChildren(QLabel))
    assert window.mark_normal.isEnabled()
    first_id=window.preview_id
    window.mark_normal.click()
    assert rows[0].status == Status.NORMAL and window.model.rowCount() == 1
    assert window.current is rows[1] and window.counts[Status.NORMAL] == 2
    assert window.counts[Status.REVIEW] == 0 and window.counts[Status.ERROR] == 1
    window.pool.waitForDone(); app.processEvents()
    window.mark_normal.click()
    assert window.proxy.rowCount() == 0 and window.current is None and window.view.photo is None
    assert not window.copy.isEnabled() and not window.mark_normal.isEnabled()
    window.preview_ready(first_id,QImage(2,2,QImage.Format.Format_RGB32),'')
    assert window.view.photo is None
    assert window.export.isEnabled()
    window.close()


def test_reclassify_last_filtered_item_clears_preview_but_preserves_other_results(tmp_path):
    app,window,rows=window_fixture(tmp_path)
    window.search.setText('첫 폴더')
    window.pool.waitForDone(); app.processEvents()
    window.mark_normal.click()
    assert window.current is None and window.proxy.rowCount() == 0 and window.model.rowCount() == 1
    assert not window.copy.isEnabled()
    window.search.clear()
    assert window.current is rows[1]
    window.close()


def test_copy_button_uses_filtered_snapshot_and_export_includes_all_non_normal(tmp_path,monkeypatch):
    app,window,rows=window_fixture(tmp_path)
    window.search.setText('첫 폴더')
    destination=tmp_path/'선택 사본'; destination.mkdir()
    monkeypatch.setattr(QFileDialog,'getExistingDirectory',lambda *a:str(destination))
    alerts = []
    monkeypatch.setattr(QMessageBox,'exec',lambda self: alerts.append(self.text()) or 0)
    window.copy.click()
    assert window.copy_progress.isVisible() and window.copy_progress.maximum() == 1
    assert not window.start.isEnabled() and not window.mark_normal.isEnabled()
    loop=QEventLoop(); timer=QTimer(); timer.setSingleShot(True)
    timer.timeout.connect(loop.quit); timer.start(5000)
    window.copy_worker.finished.connect(loop.quit)
    loop.exec(); timer.stop(); app.processEvents()
    assert window.copy_worker is None and '복사 1건' in window.state.text()
    assert window.copy_progress is None and len(alerts) == 1 and '사진 복사 완료' in alerts[0]
    assert (destination/rows[0].relative_path).exists()
    assert not (destination/rows[1].relative_path).exists()
    monkeypatch.setattr(QFileDialog,'getSaveFileName',lambda *a:(str(tmp_path/'결과'),'엑셀'))
    window.export.click()
    assert len(read_rows(tmp_path/'결과.xlsx')) == 2
    window.close()


def test_reclassification_disabled_during_scan_or_without_valid_preview(tmp_path):
    app,window,rows=window_fixture(tmp_path)
    window.set_running(True)
    window.reclassify_normal()
    assert rows[0].status == Status.REVIEW and not window.mark_normal.isEnabled()
    window.set_running(False)
    window.view.clear(); window.update_actions()
    window.reclassify_normal()
    assert rows[0].status == Status.REVIEW
    window.close()

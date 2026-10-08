import csv
from dataclasses import replace
import json
import os
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
import pytest
from PySide6.QtWidgets import QApplication, QMessageBox

from photocheck.criteria import Criteria, load_criteria, save_criteria
from photocheck.export import export_excel
from photocheck.settings_dialog import SettingsDialog
from photocheck.ui import MainWindow


@pytest.mark.parametrize("values", [{"roll_degrees":float('nan')}, {"confidence":1.1},
    {"face_height_min":True}, {"check_recapture":1}, {"nose_height_min":.8}, {"paper_area_min":.9}])
def test_invalid_settings_are_rejected(values):
    with pytest.raises(ValueError):
        replace(Criteria(), **values)


def test_settings_roundtrip_and_corrupt_file_fallback(tmp_path):
    path = tmp_path/'사용자 데이터'/'설정.json'
    custom = replace(Criteria(), roll_degrees=8, check_recapture=False)
    save_criteria(custom,path)
    assert load_criteria(path) == (custom, '')
    assert custom.identity != Criteria().identity
    path.write_text('{broken',encoding='utf-8')
    criteria,error = load_criteria(path)
    assert criteria == Criteria() and error
    assert path.read_text() == '{broken'


def test_previous_version_preferences_keep_custom_values_without_rewriting(tmp_path):
    path = tmp_path/'settings.json'
    old=replace(Criteria(),face_height_min=.31,roll_degrees=8,check_recapture=False).snapshot()
    old['version']='strict-0.5-settings-recapture-unvalidated'
    for key in ('check_aspect','aspect_width','aspect_height','aspect_tolerance'):
        old.pop(key)
    path.write_text(json.dumps({'schema':1,'criteria':old}),encoding='utf-8')
    before=path.read_bytes()
    criteria,error=load_criteria(path)
    assert not error and criteria.face_height_min==.31 and criteria.roll_degrees==8
    assert not criteria.check_recapture and criteria.check_aspect
    assert criteria.aspect_width==3 and criteria.aspect_height==4 and criteria.aspect_tolerance==.5
    assert path.read_bytes()==before


def test_percentage_preferences_migrate_without_treating_5_as_ratio_units(tmp_path):
    path=tmp_path/'settings.json'
    values=replace(Criteria(),aspect_width=3.5,aspect_height=4.5,check_aspect=False,roll_degrees=8).snapshot()
    values['version']='strict-0.6-aspect-ratio-unvalidated'
    values.pop('aspect_tolerance')
    values['aspect_tolerance_pct']=5
    path.write_text(json.dumps({'schema':1,'criteria':values}),encoding='utf-8')
    before=path.read_bytes()
    criteria,error=load_criteria(path)
    assert not error and criteria.aspect_tolerance==.5
    assert criteria.aspect_width==3.5 and criteria.aspect_height==4.5
    assert criteria.roll_degrees==8 and not criteria.check_aspect and path.read_bytes()==before


def test_dialog_ratio_tolerance_label_and_saved_value(tmp_path):
    app=QApplication.instance() or QApplication([])
    from PySide6.QtWidgets import QLabel
    dialog=SettingsDialog(Criteria(),path=tmp_path/'settings.json')
    labels=[label.text() for label in dialog.findChildren(QLabel)]
    assert '허용 오차 비율' in labels and '비율 허용 오차 (%)' not in labels
    assert any('세로−가로 0.5~1.5' in label for label in labels)
    dialog.fields['aspect_tolerance'][0].setValue(.25)
    dialog.save()
    assert load_criteria(tmp_path/'settings.json')[0].aspect_tolerance==.25
    dialog.close()


def test_failed_save_preserves_existing_preferences(tmp_path, monkeypatch):
    path = tmp_path/'settings.json'
    save_criteria(Criteria(),path)
    before = path.read_bytes()
    def denied(*args):
        raise PermissionError('locked')
    monkeypatch.setattr('photocheck.criteria.os.replace', denied)
    with pytest.raises(PermissionError):
        save_criteria(replace(Criteria(),face_height_min=.3),path)
    assert path.read_bytes() == before and not list(tmp_path.glob('.settings-*.tmp'))


def test_dialog_cancel_restore_save_and_validation(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])
    path = tmp_path/'settings.json'
    custom = replace(Criteria(),face_height_min=.3)
    dialog = SettingsDialog(custom,path=path)
    dialog.fields['face_height_min'][0].setValue(40)
    dialog.reject()
    assert dialog.criteria == custom and not path.exists()
    dialog = SettingsDialog(custom,path=path)
    dialog.populate(Criteria())
    assert dialog.fields['face_height_min'][0].value() == 25
    dialog.checks['check_recapture'].setChecked(False)
    assert not dialog.fields['paper_area_min'][0].isEnabled()
    dialog.save()
    assert dialog.result() == SettingsDialog.DialogCode.Accepted
    assert load_criteria(path)[0] == replace(Criteria(),check_recapture=False)
    dialog = SettingsDialog(custom,path=path)
    before = path.read_bytes()
    notices=[]
    monkeypatch.setattr(QMessageBox,'warning',lambda *args:notices.append(args))
    dialog.fields['nose_height_min'][0].setValue(90)
    dialog.save()
    assert notices and dialog.result() != SettingsDialog.DialogCode.Accepted and path.read_bytes() == before
    dialog.close()


def test_window_loads_preferences_and_disables_gear_while_scanning(tmp_path,monkeypatch):
    app = QApplication.instance() or QApplication([])
    monkeypatch.setenv('LOCALAPPDATA',str(tmp_path))
    custom = replace(Criteria(),roll_degrees=7)
    save_criteria(custom)
    window = MainWindow()
    assert window.criteria == custom
    assert not window.settings.icon().isNull()
    window.set_running(True)
    assert not window.settings.isEnabled()
    window.set_running(False)
    assert window.settings.isEnabled()
    assert window.reason_filter.findText('재촬영 의심') >= 0
    assert window.reason_filter.findText('사진 비율') >= 0
    assert window.reason_filter.findText('글자 포함 의심') >= 0
    window.close()


def test_excel_omits_settings_and_preserves_internal_result_settings(tmp_path):
    from openpyxl import load_workbook
    from photocheck.domain import Result
    original = Criteria()
    custom = replace(original,face_height_min=.4)
    rows = [Result(tmp_path/f'{i}.jpg',f'{i}.jpg',criteria_id=c.identity,criteria_settings=c.snapshot())
            for i,c in enumerate((original,custom))]
    export_excel(tmp_path/'result.xlsx',rows)
    book = load_workbook(tmp_path/'result.xlsx')
    values = list(book.active.values)
    output = [dict(zip(values[0],row)) for row in values[1:]]
    assert values[0] == ('파일명', '판정사유')
    assert len(output) == 2 and book.active.max_column == 2
    assert rows[0].criteria_settings['face_height_min'] == .25
    assert rows[1].criteria_settings['face_height_min'] == .4
    assert rows[0].criteria_id != rows[1].criteria_id
    book.close()
